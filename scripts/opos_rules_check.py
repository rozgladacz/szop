"""Semantyczny drift-check normatywnego DOCX, publikowanego PDF i YAML OPOS."""

from __future__ import annotations

import re
import sys
import unicodedata
from decimal import Decimal
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app.services.opos_rules import load_opos_ruleset  # noqa: E402
from app.services.opos_rules.models import OposRuleset  # noqa: E402


SPRITE_PATH = ROOT_DIR / "app" / "static" / "icons" / "opos.svg"


def normalize(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", text).lower()
    return " ".join(re.sub(r"[^\w]+", " ", normalized).split())


def extract_docx_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def extract_docx_abilities(path: Path) -> dict[str, set[str]]:
    """Read the three normative ability lists independently of YAML."""
    document = Document(path)
    headings = {
        "zdolności pasywne": "passive",
        "zdolności specjalne": "special",
        "zdolności ataku": "weapon",
        "zdolności broni": "weapon",
    }
    result = {category: set() for category in headings.values()}
    category: str | None = None
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        normalized = normalize(value)
        matched_heading = next(
            (heading for heading in headings if normalized.startswith(heading)), None
        )
        if matched_heading is not None:
            category = headings[matched_heading]
            continue
        if normalized.startswith("koszt oddziału"):
            break
        if category is None or ":" not in value:
            continue
        name = value.split(":", maxsplit=1)[0].strip()
        if normalize(name).startswith("aura"):
            name = "Aura"
        result[category].add(name)
    return result


def extract_pdf_text(path: Path) -> tuple[str, int]:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)


def _semantic_document_errors(
    text: str, *, label: str, ruleset: OposRuleset
) -> list[str]:
    normalized = normalize(text)
    errors: list[str] = []
    for ability in ruleset.abilities:
        if normalize(ability.name) not in normalized:
            errors.append(f"{label}: brak zdolności {ability.name}")

    critical_slugs = (
        "fast",
        "steadfast",
        "patient",
        "breakthrough",
        "airplane",
        "clumsy",
        "deadly",
        "transport",
        "area",
        "double",
        "charge",
        "prepared",
    )
    by_slug = ruleset.abilities_by_slug
    for slug in critical_slugs:
        ability = by_slug[slug]
        if normalize(ability.description) not in normalized:
            errors.append(f"{label}: opis {ability.name} różni się od YAML")

    formula_facts = {
        "modyfikator życia": (
            "Modyfikator życia wynosi 1+0,01 * życie"
        ),
        "modyfikator zbroi": "Modyfikator zbroi: 0,09x^2-0,19x+1",
        "modyfikator siły": "Modyfikator siły: -0,04x^2 +0,5x+1",
        "koszt broni": (
            "Koszt broni wynosi: liczba kości * 6 * modyfikator zasięgu * "
            "modyfikator siły * modyfikatory zdolności za każdą broń"
        ),
        "najdroższy profil liczony dwukrotnie": (
            "Koszt najdroższego profilu policz dwukrotnie"
        ),
        "zasięgi 0,5/0,6/1": (
            "Zasięg Wręcz Krótki Długi Modyfikator 0,5 0,6 1"
        ),
        "Samolot -1": "Samolot: odlicz 1 od kosztu zdolności",
        "Samolot 0/0/1,6": "Modyfikator zasięgu: 0/0/1,6",
        "Niezgrabny -1": "Niezgrabny: odlicz 1 od kosztu zdolności",
        "Zabójczy ×4": "Zabójczy: *4",
        "Transport +2": "Transport: dolicz 2 do kosztu zdolności",
        "Podwójny — sukces także remisem": (
            "Podwójny: Każdy sukces liczy się również jako remis"
        ),
        "Aura dla profili ataku": "lub ich profile ataku",
        "Szarża ×1,4": "Szarża: *1,4",
        "Przygotowanie ×1,4": "Przygotowanie: *1,4",
    }
    area_small = by_slug["area"].small_battle_description
    if area_small and normalize(area_small) not in normalized:
        errors.append(f"{label}: brak wariantowego opisu Obszarowej")
    for fact_name, fragment in formula_facts.items():
        if normalize(fragment) not in normalized:
            errors.append(f"{label}: brak semantyki „{fact_name}”")
    return errors


def _yaml_contract_errors(ruleset: OposRuleset) -> list[str]:
    errors: list[str] = []
    if ruleset.version != "1.2.0":
        errors.append("YAML: hotfix musi mieć wersję 1.2.0")
    expected_stats = {
        "defense": (Decimal("3"), Decimal("4"), Decimal("5")),
        "toughness": tuple(Decimal(value) for value in (2, 4, 6, 12, 18, 24)),
        "strength": (Decimal("0"), Decimal("1"), Decimal("2")),
    }
    for field, expected in expected_stats.items():
        if tuple(getattr(ruleset.standard_stats, field)) != expected:
            errors.append(f"YAML: niepoprawna standardowa lista {field}")

    expected_ranges = {
        "melee": Decimal("0.5"),
        "short": Decimal("0.6"),
        "long": Decimal("1"),
    }
    for slug, multiplier in expected_ranges.items():
        if ruleset.ranges[slug].multiplier != multiplier:
            errors.append(f"YAML: niepoprawny mnożnik zasięgu {slug}")

    by_slug = ruleset.abilities_by_slug
    if by_slug["airplane"].effects.ability_cost_delta != Decimal("-1"):
        errors.append("YAML: Samolot musi kosztować -1")
    expected_airplane_ranges = {
        "melee": Decimal("0"),
        "short": Decimal("0"),
        "long": Decimal("1.6"),
    }
    if by_slug["airplane"].effects.profile_multipliers != expected_airplane_ranges:
        errors.append("YAML: Samolot musi mieć mnożniki profili 0/0/1,6")
    if by_slug["airplane"].aura_eligible:
        errors.append("YAML: Samolot nie może być celem Aury")
    if by_slug["clumsy"].effects.ability_cost_delta != Decimal("-1"):
        errors.append("YAML: Niezgrabny musi kosztować -1")
    if by_slug["clumsy"].aura_eligible:
        errors.append("YAML: Niezgrabny nie może być celem Aury")
    if "guardian" in by_slug:
        errors.append("YAML: Strażnik musi zostać usunięty")
    if not by_slug["area"].small_battle_description:
        errors.append("YAML: Obszarowa wymaga opisu dla małej bitwy")
    if by_slug["deadly"].effects.weapon_multiplier != Decimal("4"):
        errors.append("YAML: Zabójczy musi mieć mnożnik ×4")
    if by_slug["transport"].effects.ability_cost_delta != Decimal("2"):
        errors.append("YAML: Transport musi kosztować +2")
    if by_slug["double"].effects.weapon_multiplier != Decimal("1.5"):
        errors.append("YAML: Podwójny musi mieć mnożnik ×1,5")

    expected_weapon_contracts = {
        "charge": (("melee",), Decimal("1.4")),
        "prepared": (("short", "long"), Decimal("1.4")),
    }
    for slug, (ranges, multiplier) in expected_weapon_contracts.items():
        ability = by_slug[slug]
        if ability.category != "weapon":
            errors.append(f"YAML: {ability.name} musi być zdolnością broni")
        if ability.allowed_ranges != ranges:
            errors.append(f"YAML: {ability.name} ma niepoprawne zasięgi")
        if ability.effects.weapon_multiplier != multiplier:
            errors.append(f"YAML: {ability.name} musi mieć mnożnik ×1,4")

    weapon_slugs = {
        ability.slug for ability in ruleset.abilities_of("weapon")
    }
    aura_weapon_slugs = {
        ability.slug
        for ability in ruleset.aura_targets
        if ability.category == "weapon"
    }
    if aura_weapon_slugs != weapon_slugs:
        errors.append("YAML: Aura musi dopuszczać wszystkie zdolności broni")

    expected_formula = {
        "base_constant": Decimal("6"),
        "defense_quadratic": Decimal("0.09"),
        "defense_linear": Decimal("-0.19"),
        "weapon_factor": Decimal("6"),
        "strength_quadratic": Decimal("-0.04"),
        "strength_linear": Decimal("0.5"),
        "toughness_modifier_per_point": Decimal("0.01"),
    }
    for field, expected in expected_formula.items():
        if getattr(ruleset.formula, field) != expected:
            errors.append(f"YAML: niepoprawny współczynnik formuły {field}")
    return errors


def _icon_errors(ruleset: OposRuleset, sprite_path: Path) -> list[str]:
    sprite = sprite_path.read_text(encoding="utf-8")
    available = set(re.findall(r'id="icon-([a-z0-9-]+)"', sprite))
    required = set(ruleset.stat_icons.values())
    required.update(item.icon for item in ruleset.ranges.values())
    required.update(item.icon for item in ruleset.abilities)
    missing = sorted(required - available)
    return [f"SVG: brak ikony {name}" for name in missing]


def run_checks(
    *,
    ruleset: OposRuleset | None = None,
    docx_path: Path | None = None,
    pdf_path: Path | None = None,
    sprite_path: Path = SPRITE_PATH,
) -> list[str]:
    active_ruleset = ruleset or load_opos_ruleset()
    normative = docx_path or ROOT_DIR / active_ruleset.sources.normative
    published = pdf_path or ROOT_DIR / active_ruleset.sources.published

    docx_text = extract_docx_text(normative)
    docx_abilities = extract_docx_abilities(normative)
    pdf_text, pdf_pages = extract_pdf_text(published)
    errors = _yaml_contract_errors(active_ruleset)
    errors.extend(
        _semantic_document_errors(docx_text, label="DOCX", ruleset=active_ruleset)
    )
    for category, names in docx_abilities.items():
        yaml_names = {ability.name for ability in active_ruleset.abilities_of(category)}
        missing_in_yaml = sorted(names - yaml_names)
        missing_in_docx = sorted(yaml_names - names)
        if missing_in_yaml:
            errors.append(
                f"DOCX/YAML: dodatkowe zdolności {category}: "
                + ", ".join(missing_in_yaml)
            )
        if missing_in_docx:
            errors.append(
                f"DOCX/YAML: brak zdolności {category}: "
                + ", ".join(missing_in_docx)
            )
    errors.extend(
        _semantic_document_errors(pdf_text, label="PDF", ruleset=active_ruleset)
    )
    if pdf_pages != 2:
        errors.append(f"PDF: oczekiwano 2 stron, znaleziono {pdf_pages}")
    errors.extend(_icon_errors(active_ruleset, sprite_path))
    return errors


def main() -> int:
    errors = run_checks()
    if errors:
        print("OPOS rules check: FAILED")
        for error in errors:
            print(f"- {error}")
        return 1
    ruleset = load_opos_ruleset()
    print(
        "OPOS rules check: OK — "
        f"{len(ruleset.abilities)} zdolności, komplet ikon, DOCX/PDF/YAML zgodne"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
